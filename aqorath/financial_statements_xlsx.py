"""Pure in-memory XLSX renderer for formal financial statements.

This renderer consumes only a previously built ``FinancialStatementsBundle``.
The bundle's formal views remain authoritative for section totals and result;
this module performs presentation structure and line-sign display only. It never
queries storage/catalog/runtime, rebuilds statements, uses formulas, converts
money to binary floats, or writes files.
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .financial_statements import FinancialStatementsBundle
from . import reporting_xlsx as _xlsx_infra


_HEADERS = ("Código", "Cuenta", "Subtipo", "Importe")


def _configure_workbook(workbook):
    fixed_timestamp = datetime(2000, 1, 1, 0, 0, 0)
    workbook.properties.created = fixed_timestamp
    workbook.properties.modified = fixed_timestamp
    workbook.properties.creator = "Aqorath"
    workbook.properties.lastModifiedBy = "Aqorath"
    workbook.properties.title = "Aqorath Estados Financieros"


def _configure_sheet(worksheet, title, as_of):
    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = title
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center")

    worksheet["A2"] = "Al"
    worksheet["B2"] = as_of

    for column, header in enumerate(_HEADERS, start=1):
        cell = worksheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True)

    worksheet.freeze_panes = "A5"
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 42
    worksheet.column_dimensions["C"].width = 28
    worksheet.column_dimensions["D"].width = 22


def _write_section_label(worksheet, row_index, label):
    worksheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
    return row_index + 1


def _write_detail_row(worksheet, row_index, code, name, subtype, amount):
    worksheet.cell(row=row_index, column=1, value=code)
    worksheet.cell(row=row_index, column=2, value=name)
    worksheet.cell(row=row_index, column=3, value=subtype)
    worksheet.cell(row=row_index, column=4, value=str(amount))
    return row_index + 1


def _write_total_row(worksheet, row_index, label, amount):
    label_cell = worksheet.cell(row=row_index, column=2, value=label)
    amount_cell = worksheet.cell(row=row_index, column=4, value=str(amount))
    label_cell.font = Font(bold=True)
    amount_cell.font = Font(bold=True)
    return row_index + 1


def _write_income_statement(worksheet, view):
    _configure_sheet(worksheet, "Estado de Resultados", view.as_of)
    row_index = 5

    row_index = _write_section_label(worksheet, row_index, "INGRESOS")
    for line in view.income.lines:
        row_index = _write_detail_row(
            worksheet,
            row_index,
            line.account_code,
            line.account_name,
            line.account_subtype,
            -line.ledger_balance,
        )
    row_index = _write_total_row(worksheet, row_index, "Total ingresos", view.income.total)

    row_index += 1
    row_index = _write_section_label(worksheet, row_index, "COSTOS")
    for line in view.costs.lines:
        row_index = _write_detail_row(
            worksheet,
            row_index,
            line.account_code,
            line.account_name,
            line.account_subtype,
            line.ledger_balance,
        )
    row_index = _write_total_row(worksheet, row_index, "Total costos", view.costs.total)

    row_index += 1
    row_index = _write_section_label(worksheet, row_index, "GASTOS")
    for line in view.expenses.lines:
        row_index = _write_detail_row(
            worksheet,
            row_index,
            line.account_code,
            line.account_name,
            line.account_subtype,
            line.ledger_balance,
        )
    row_index = _write_total_row(worksheet, row_index, "Total gastos", view.expenses.total)

    row_index += 1
    _write_total_row(worksheet, row_index, "Resultado del ejercicio", view.result)


def _write_balance_section(worksheet, row_index, label, section, total_label):
    row_index = _write_section_label(worksheet, row_index, label)
    for line in section.lines:
        row_index = _write_detail_row(
            worksheet,
            row_index,
            line.account_code,
            line.account_name,
            line.account_subtype,
            line.amount,
        )
    return _write_total_row(worksheet, row_index, total_label, section.total)


def _write_balance_sheet(worksheet, view):
    _configure_sheet(worksheet, "Balance General", view.as_of)
    row_index = 5

    row_index = _write_balance_section(
        worksheet, row_index, "ACTIVO", view.assets, "Total activo"
    )
    row_index += 1
    row_index = _write_balance_section(
        worksheet, row_index, "PASIVO", view.liabilities, "Total pasivo"
    )
    row_index += 1
    row_index = _write_balance_section(
        worksheet,
        row_index,
        "PATRIMONIO",
        view.recorded_equity,
        "Total patrimonio registrado",
    )

    row_index = _write_total_row(
        worksheet, row_index, "Resultado del ejercicio", view.current_result
    )
    row_index = _write_total_row(
        worksheet, row_index, "Total patrimonio", view.total_equity
    )
    _write_total_row(
        worksheet,
        row_index,
        "Total pasivo y patrimonio",
        view.liabilities_and_equity,
    )


def render_financial_statements_xlsx(bundle):
    """Serialize one formal statements bundle to an in-memory two-sheet XLSX."""
    if not isinstance(bundle, FinancialStatementsBundle):
        raise TypeError("bundle must be a FinancialStatementsBundle")

    workbook = Workbook()
    _configure_workbook(workbook)

    income_sheet = workbook.active
    income_sheet.title = "Estado de Resultados"
    balance_sheet = workbook.create_sheet("Balance General")

    _write_income_statement(income_sheet, bundle.income_statement)
    _write_balance_sheet(balance_sheet, bundle.balance_sheet)

    data = _xlsx_infra._save_workbook_in_memory(workbook)
    workbook.close()
    return data
