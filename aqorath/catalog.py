from pathlib import Path
import csv
from typing import Dict, Optional
from openpyxl import load_workbook

from sqlmodel import select
from .storage import get_session
from .models import Account, AppConfig

def _read_xlsx(path: Path) -> Dict[str, Dict[str, str]]:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb.active
    catalog = {}
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx == 1:
            continue
        code = row[0]
        if not code:
            continue
        code = str(code).strip()
        catalog[code] = {
            "name_osc": (str(row[1]).strip() if row[1] else ""),
            "name_comercial": (str(row[2]).strip() if row[2] else ""),
            "tipo": (str(row[3]).strip() if row[3] else ""),
            "subtipo": (str(row[4]).strip() if row[4] else ""),
            "naturaleza": (str(row[5]).strip() if row[5] else ""),
            "descripcion": (str(row[6]).strip() if row[6] else ""),
        }
    return catalog

def _read_csv(path: Path) -> Dict[str, Dict[str, str]]:
    catalog = {}
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        for r_idx, row in enumerate(reader, start=1):
            if r_idx == 1:
                continue
            if not row:
                continue
            code = row[0].strip()
            if not code:
                continue
            def col(i):
                return row[i].strip() if i < len(row) and row[i] is not None else ""
            catalog[code] = {
                "name_osc": col(1),
                "name_comercial": col(2),
                "tipo": col(3),
                "subtipo": col(4),
                "naturaleza": col(5),
                "descripcion": col(6),
            }
    return catalog

def load_catalog(path: Optional[Path] = None) -> Dict[str, Dict[str, str]]:
    path = Path(path) if path else Path("assets/catalogo.csv")
    if not path.exists():
        alt = Path("assets/Catálogo.xlsx")
        if alt.exists():
            path = alt
        else:
            return {}
    if path.suffix.lower() in (".xls", ".xlsx"):
        return _read_xlsx(path)
    return _read_csv(path)

def get_catalog_entry(code: str, path: Optional[Path] = None) -> Dict[str, str]:
    return load_catalog(path).get(str(code), {})

def resolve_account_by_code(session, code: str, prefer: Optional[str] = None):
    """
    Resolver un Account por código, valorando duplicados:
    - prefer: 'osc' | 'comercial' | None
    Lógica:
      1) si existe AppConfig 'account.<code>.canonical_id.{prefer}' -> devuelve esa fila
      2) luego si existe 'account.<code>.canonical_id' -> devuelve esa fila
      3) consulta rows con ese code; si hay una sola -> return
      4) si hay múltiples: intentar match por catálogo name_{prefer}, luego AppConfig name_{prefer}, luego fallback first()
    """
    # prefer-specific canonical
    if prefer:
        key_pref = f"account.{code}.canonical_id.{prefer}"
        cfg_pref = session.exec(select(AppConfig).where(AppConfig.key == key_pref)).one_or_none()
        if cfg_pref and cfg_pref.value:
            try:
                aid = int(cfg_pref.value)
                acc = session.exec(select(Account).where(Account.id == aid)).one_or_none()
                if acc:
                    return acc
            except Exception:
                pass

    # general canonical_id
    cfg_key = f"account.{code}.canonical_id"
    cfg = session.exec(select(AppConfig).where(AppConfig.key == cfg_key)).one_or_none()
    if cfg and cfg.value:
        try:
            aid = int(cfg.value)
            acc = session.exec(select(Account).where(Account.id == aid)).one_or_none()
            if acc:
                return acc
        except Exception:
            pass

    rows = session.exec(select(Account).where(Account.code == str(code))).all()
    if not rows:
        return None
    if len(rows) == 1:
        return rows[0]

    # multiples: usar catálogo
    catalog = load_catalog()
    entry = catalog.get(str(code), {})
    preferred_name = ""
    if prefer == "osc":
        preferred_name = entry.get("name_osc") or ""
    elif prefer == "comercial":
        preferred_name = entry.get("name_comercial") or ""

    def norm(s: str) -> str:
        return (s or "").strip().lower()

    if preferred_name:
        for r in rows:
            if norm(getattr(r, "name", "")) == norm(preferred_name):
                return r

    # fallback: comprobar AppConfig name entries por prefer
    name_pref = session.exec(select(AppConfig).where(AppConfig.key == f"account.{code}.name_{prefer}")).one_or_none() if prefer else None
    if name_pref and name_pref.value:
        for r in rows:
            if norm(getattr(r, "name", "")) == norm(name_pref.value):
                return r

    # check both name_comercial/name_osc stored in AppConfig generally
    name_com = session.exec(select(AppConfig).where(AppConfig.key == f"account.{code}.name_comercial")).one_or_none()
    name_osc = session.exec(select(AppConfig).where(AppConfig.key == f"account.{code}.name_osc")).one_or_none()
    if name_com and name_com.value:
        for r in rows:
            if norm(getattr(r, "name", "")) == norm(name_com.value):
                return r
    if name_osc and name_osc.value:
        for r in rows:
            if norm(getattr(r, "name", "")) == norm(name_osc.value):
                return r

    # nada coincide: devolver la primera
    return rows[0]