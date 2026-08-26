"""Pure in-memory XLSX renderer for immutable financial-report snapshots.

This module is presentation-only. It consumes a previously built
``FinancialReportSnapshot`` and serializes it to XLSX bytes without querying
storage, loading catalog data, recalculating accounting, or writing files.
Monetary values are stored as exact decimal text rather than binary floats or
spreadsheet formulas.
"""

from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Font
from openpyxl.writer.excel import ExcelWriter
from openpyxl.worksheet._writer import WorksheetWriter

from .reporting import FinancialReportSnapshot


_HEADERS = (
    "Account Code",
    "Account Name",
    "Type",
    "Subtype",
    "Nature",
    "Ledger Balance",
    "Normal Balance",
)


class _InMemoryExcelWriter(ExcelWriter):
    """ExcelWriter variant that keeps worksheet XML in memory.

    openpyxl's default ``WorksheetWriter`` creates a temporary file even when
    the final workbook target is a ``BytesIO`` object. Reporting renderers must
    not touch the filesystem, so worksheet XML is written to ``BytesIO`` and
    inserted into the XLSX archive with ``writestr`` instead.
    """

    def write_worksheet(self, worksheet):
        worksheet._drawing = SpreadsheetDrawing()
        worksheet._drawing.charts = worksheet._charts
        worksheet._drawing.images = worksheet._images

        worksheet_xml = BytesIO()
        writer = WorksheetWriter(worksheet, worksheet_xml)
        writer.write()

        worksheet._rels = writer._rels
        self._archive.writestr(
            worksheet.path[1:],
            worksheet_xml.getvalue(),
        )
        self.manifest.append(worksheet)


def _save_workbook_in_memory(workbook):
    output = BytesIO()
    archive = ZipFile(
        output,
        mode="w",
        compression=ZIP_DEFLATED,
        allowZip64=True,
    )
    writer = _InMemoryExcelWriter(workbook, archive)
    writer.save()
    return output.getvalue()


def render_financial_report_xlsx(snapshot):
    """Serialize one ``FinancialReportSnapshot`` to in-memory XLSX bytes."""
    if not isinstance(snapshot, FinancialReportSnapshot):
        raise TypeError("snapshot must be a FinancialReportSnapshot")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Financial Report"

    # Keep package metadata stable and independent of wall-clock time.
    fixed_timestamp = datetime(2000, 1, 1, 0, 0, 0)
    workbook.properties.created = fixed_timestamp
    workbook.properties.modified = fixed_timestamp
    workbook.properties.creator = "Aqorath"
    workbook.properties.lastModifiedBy = "Aqorath"
    workbook.properties.title = "Aqorath Financial Report"

    worksheet.merge_cells("A1:G1")
    worksheet["A1"] = "Aqorath Financial Report"
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center")

    worksheet["A2"] = "As of"
    worksheet["B2"] = snapshot.as_of

    for column, header in enumerate(_HEADERS, start=1):
        cell = worksheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True)

    row_index = 5
    for line in snapshot.lines:
        values = (
            line.account_code,
            line.account_name,
            line.account_type,
            line.account_subtype,
            line.nature,
            str(line.ledger_balance),
            str(line.normal_balance),
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(row=row_index, column=column, value=value)
        row_index += 1

    # Separate detail from snapshot totals while preserving total order verbatim.
    row_index += 1
    for total in snapshot.totals:
        worksheet.cell(row=row_index, column=1, value="TOTAL")
        worksheet.cell(row=row_index, column=3, value=total.account_type)
        worksheet.cell(row=row_index, column=7, value=str(total.amount))
        worksheet.cell(row=row_index, column=1).font = Font(bold=True)
        worksheet.cell(row=row_index, column=3).font = Font(bold=True)
        worksheet.cell(row=row_index, column=7).font = Font(bold=True)
        row_index += 1

    worksheet.cell(row=row_index, column=1, value="RESULT")
    worksheet.cell(row=row_index, column=2, value="Resultado del ejercicio")
    worksheet.cell(row=row_index, column=7, value=str(snapshot.result))
    for column in (1, 2, 7):
        worksheet.cell(row=row_index, column=column).font = Font(bold=True)

    worksheet.freeze_panes = "A5"
    if snapshot.lines:
        worksheet.auto_filter.ref = f"A4:G{4 + len(snapshot.lines)}"

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 38
    worksheet.column_dimensions["C"].width = 16
    worksheet.column_dimensions["D"].width = 28
    worksheet.column_dimensions["E"].width = 14
    worksheet.column_dimensions["F"].width = 22
    worksheet.column_dimensions["G"].width = 22

    data = _save_workbook_in_memory(workbook)
    workbook.close()
    return data
