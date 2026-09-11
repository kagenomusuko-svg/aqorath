"""Presentation-only XLSX/PDF rendering for AQR-013 governed report products.

Renderers consume already-built semantic content.  They never query storage, recompute
balances, apply fiscal rules or persist output.  Exact Decimal values are serialized as
text, matching the established Aqorath reporting exporters.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from . import reporting_xlsx as _xlsx_infra


_QUERY_TRIAL = "period_trial_balance"
_QUERY_INCOME = "period_income_statement"
_QUERY_BALANCE = "as_of_balance_sheet"


def _configure_workbook(workbook, title):
    fixed = datetime(2000, 1, 1, 0, 0, 0)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.properties.creator = "Aqorath"
    workbook.properties.lastModifiedBy = "Aqorath"
    workbook.properties.title = title


def _sheet_header(ws, title, request, headers):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Desde"
    ws["B2"] = request.from_date.isoformat()
    ws["C2"] = "Hasta"
    ws["D2"] = request.to_date.isoformat()
    for idx, value in enumerate(headers, 1):
        cell = ws.cell(row=4, column=idx, value=value)
        cell.font = Font(bold=True)
    ws.freeze_panes = "A5"


def _write_trial_xlsx(ws, report):
    _sheet_header(
        ws,
        report.definition.name,
        report.request,
        ("Código", "Cuenta", "Apertura", "Debe", "Haber", "Cierre"),
    )
    row = 5
    for line in report.content.lines:
        values = (
            line.account_code,
            line.account_name,
            str(line.opening_balance),
            str(line.debits),
            str(line.credits),
            str(line.closing_balance),
        )
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
        row += 1
    ws.cell(row=row, column=2, value="Totales del período").font = Font(bold=True)
    ws.cell(row=row, column=4, value=str(report.content.total_debits)).font = Font(bold=True)
    ws.cell(row=row, column=5, value=str(report.content.total_credits)).font = Font(bold=True)


def _write_income_xlsx(ws, report):
    _sheet_header(
        ws,
        report.definition.name,
        report.request,
        ("Código", "Cuenta", "Subtipo", "Importe"),
    )
    row = 5
    for label, section in (
        ("INGRESOS", report.content.income),
        ("COSTOS", report.content.costs),
        ("GASTOS", report.content.expenses),
    ):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        row += 1
        for line in section.lines:
            amount = -line.ledger_balance if line.account_type == "Ingreso" else line.ledger_balance
            values = (
                line.account_code,
                line.account_name,
                line.account_subtype,
                str(amount),
            )
            for col, value in enumerate(values, 1):
                ws.cell(row=row, column=col, value=value)
            row += 1
        ws.cell(row=row, column=2, value=f"Total {label.lower()}").font = Font(bold=True)
        ws.cell(row=row, column=4, value=str(section.total)).font = Font(bold=True)
        row += 2
    ws.cell(row=row, column=2, value="Resultado del período").font = Font(bold=True)
    ws.cell(row=row, column=4, value=str(report.content.result)).font = Font(bold=True)


def _write_balance_xlsx(ws, report):
    _sheet_header(
        ws,
        report.definition.name,
        report.request,
        ("Código", "Cuenta", "Subtipo", "Importe"),
    )
    row = 5
    for label, section in (
        ("ACTIVO", report.content.assets),
        ("PASIVO", report.content.liabilities),
        ("PATRIMONIO REGISTRADO", report.content.recorded_equity),
    ):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        row += 1
        for line in section.lines:
            values = (
                line.account_code,
                line.account_name,
                line.account_subtype,
                str(line.amount),
            )
            for col, value in enumerate(values, 1):
                ws.cell(row=row, column=col, value=value)
            row += 1
        ws.cell(row=row, column=2, value=f"Total {label.lower()}").font = Font(bold=True)
        ws.cell(row=row, column=4, value=str(section.total)).font = Font(bold=True)
        row += 2
    for label, amount in (
        ("Resultado del período", report.content.current_result),
        ("Total patrimonio", report.content.total_equity),
        ("Total pasivo y patrimonio", report.content.liabilities_and_equity),
    ):
        ws.cell(row=row, column=2, value=label).font = Font(bold=True)
        ws.cell(row=row, column=4, value=str(amount)).font = Font(bold=True)
        row += 1


def _write_report_xlsx(ws, report):
    key = report.definition.query_template_id
    if key == _QUERY_TRIAL:
        _write_trial_xlsx(ws, report)
    elif key == _QUERY_INCOME:
        _write_income_xlsx(ws, report)
    elif key == _QUERY_BALANCE:
        _write_balance_xlsx(ws, report)
    else:
        raise ValueError(f"unsupported XLSX report query: {key}")


def render_report_xlsx(report):
    workbook = Workbook()
    _configure_workbook(workbook, f"Aqorath - {report.definition.name}")
    ws = workbook.active
    ws.title = report.definition.name[:31]
    _write_report_xlsx(ws, report)
    data = _xlsx_infra._save_workbook_in_memory(workbook)
    workbook.close()
    return data


def render_package_xlsx(package_result):
    workbook = Workbook()
    _configure_workbook(workbook, f"Aqorath - {package_result.package.name}")
    for index, report in enumerate(package_result.reports):
        ws = workbook.active if index == 0 else workbook.create_sheet()
        ws.title = report.definition.name[:31]
        _write_report_xlsx(ws, report)
    data = _xlsx_infra._save_workbook_in_memory(workbook)
    workbook.close()
    return data


class _PdfWriter:
    def __init__(self, pdf):
        self.pdf = pdf
        self.y = 742
        self.title = ""
        self.request = None
        self.headers = ()

    def start(self, title, request, headers):
        self.pdf.showPage() if self.title else None
        self.title = title
        self.request = request
        self.headers = headers
        self.y = 742
        self._header()

    def _header(self):
        self.pdf.setFont("Helvetica-Bold", 14)
        self.pdf.drawString(45, self.y, self.title)
        self.y -= 20
        self.pdf.setFont("Helvetica", 9)
        self.pdf.drawString(
            45,
            self.y,
            f"Desde: {self.request.from_date.isoformat()}  Hasta: {self.request.to_date.isoformat()}",
        )
        self.y -= 20
        self.pdf.setFont("Helvetica-Bold", 8)
        positions = (45, 100, 300, 390, 455, 520)
        for index, header in enumerate(self.headers):
            self.pdf.drawString(positions[index], self.y, header)
        self.y -= 14

    def ensure(self, rows=1):
        if self.y - (14 * rows) >= 50:
            return
        self.pdf.showPage()
        self.y = 742
        self._header()

    def section(self, label):
        self.ensure(2)
        self.pdf.setFont("Helvetica-Bold", 9)
        self.pdf.drawString(45, self.y, label)
        self.y -= 14

    def row(self, values):
        self.ensure()
        positions = (45, 100, 300, 390, 455, 520)
        self.pdf.setFont("Helvetica", 8)
        for index, value in enumerate(values):
            text = str(value)
            if index == 1:
                text = text[:34]
            self.pdf.drawString(positions[index], self.y, text)
        self.y -= 14

    def total(self, label, amount):
        self.ensure()
        self.pdf.setFont("Helvetica-Bold", 8)
        self.pdf.drawString(100, self.y, label)
        self.pdf.drawRightString(565, self.y, str(amount))
        self.y -= 14


def _write_report_pdf(writer, report):
    key = report.definition.query_template_id
    if key == _QUERY_TRIAL:
        writer.start(
            report.definition.name,
            report.request,
            ("Código", "Cuenta", "Apertura", "Debe", "Haber", "Cierre"),
        )
        for line in report.content.lines:
            writer.row((
                line.account_code,
                line.account_name,
                line.opening_balance,
                line.debits,
                line.credits,
                line.closing_balance,
            ))
        writer.total("Debe del período", report.content.total_debits)
        writer.total("Haber del período", report.content.total_credits)
        return
    if key == _QUERY_INCOME:
        writer.start(
            report.definition.name,
            report.request,
            ("Código", "Cuenta", "Subtipo", "Importe"),
        )
        for label, section in (
            ("Ingresos", report.content.income),
            ("Costos", report.content.costs),
            ("Gastos", report.content.expenses),
        ):
            writer.section(label)
            for line in section.lines:
                amount = -line.ledger_balance if line.account_type == "Ingreso" else line.ledger_balance
                writer.row((line.account_code, line.account_name, line.account_subtype, amount))
            writer.total(f"Total {label.lower()}", section.total)
        writer.total("Resultado del período", report.content.result)
        return
    if key == _QUERY_BALANCE:
        writer.start(
            report.definition.name,
            report.request,
            ("Código", "Cuenta", "Subtipo", "Importe"),
        )
        for label, section in (
            ("Activo", report.content.assets),
            ("Pasivo", report.content.liabilities),
            ("Patrimonio registrado", report.content.recorded_equity),
        ):
            writer.section(label)
            for line in section.lines:
                writer.row((line.account_code, line.account_name, line.account_subtype, line.amount))
            writer.total(f"Total {label.lower()}", section.total)
        writer.total("Resultado del período", report.content.current_result)
        writer.total("Total patrimonio", report.content.total_equity)
        writer.total("Total pasivo y patrimonio", report.content.liabilities_and_equity)
        return
    raise ValueError(f"unsupported PDF report query: {key}")


def _render_pdf(reports, title):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter, pageCompression=0, invariant=1)
    pdf.setTitle(title)
    pdf.setAuthor("Aqorath")
    writer = _PdfWriter(pdf)
    for report in reports:
        _write_report_pdf(writer, report)
    pdf.save()
    return buffer.getvalue()


def render_report_pdf(report):
    return _render_pdf((report,), f"Aqorath - {report.definition.name}")


def render_package_pdf(package_result):
    return _render_pdf(package_result.reports, f"Aqorath - {package_result.package.name}")
