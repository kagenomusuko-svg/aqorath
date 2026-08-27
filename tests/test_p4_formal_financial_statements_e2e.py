"""Phase 4T — end-to-end integrity for canonical formal financial-statement exports."""

import sqlite3
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook


def _catalog():
    return {
        "1101": {"tipo": "Activo", "subtipo": "Circulante", "naturaleza": "Deudora"},
        "1205": {"tipo": "Activo", "subtipo": "No circulante", "naturaleza": "Acreedora"},
        "2101": {"tipo": "Pasivo", "subtipo": "Corto plazo", "naturaleza": "Acreedora"},
        "3101": {"tipo": "Patrimonio", "subtipo": "Capital", "naturaleza": "Acreedora"},
        "4201": {"tipo": "Ingreso", "subtipo": "Operacion", "naturaleza": "Acreedora"},
        "4299": {"tipo": "Ingreso", "subtipo": "Operacion", "naturaleza": "Deudora"},
        "5101": {"tipo": "Costo", "subtipo": "Operacion", "naturaleza": "Deudora"},
        "6101": {"tipo": "Gasto", "subtipo": "Administracion", "naturaleza": "Deudora"},
        "6199": {"tipo": "Gasto", "subtipo": "Administracion", "naturaleza": "Acreedora"},
    }


@pytest.fixture
def canonical_formal_db(tmp_path, monkeypatch):
    import aqorath.accounting_rules as accounting_rules

    db = tmp_path / "formal-statements-e2e.db"
    conn = sqlite3.connect(str(db))
    try:
        conn.executescript(
            """
            CREATE TABLE account (
                id INTEGER PRIMARY KEY,
                code TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL,
                nature TEXT NOT NULL,
                origin TEXT DEFAULT 'canonical',
                parent_id INTEGER
            );
            CREATE TABLE journalentry (
                id INTEGER PRIMARY KEY,
                date TEXT NOT NULL,
                concept TEXT
            );
            CREATE TABLE journalline (
                id INTEGER PRIMARY KEY,
                entry_id INTEGER,
                account_code TEXT,
                account_id INTEGER,
                debit TEXT NOT NULL DEFAULT '0',
                credit TEXT NOT NULL DEFAULT '0'
            );

            INSERT INTO account(id, code, name, nature) VALUES
                (1, '1101', 'Bancos', 'DEBIT'),
                (2, '1205', 'Depreciacion acumulada', 'CREDIT'),
                (3, '2101', 'Proveedores', 'CREDIT'),
                (4, '3101', 'Capital social', 'CREDIT'),
                (5, '4201', 'Ventas', 'CREDIT'),
                (6, '4299', 'Devoluciones sobre ventas', 'DEBIT'),
                (7, '5101', 'Costo de ventas', 'DEBIT'),
                (8, '6101', 'Servicios', 'DEBIT'),
                (9, '6199', 'Recuperacion de gasto', 'CREDIT');

            INSERT INTO journalentry(id, date, concept)
            VALUES (1, '2026-07-15', 'estado formal cerrado');

            INSERT INTO journalline(entry_id, account_code, account_id, debit, credit) VALUES
                (1, '1101', 1, '1000.0000', '0'),
                (1, '1205', 2, '0', '200.0000'),
                (1, '2101', 3, '0', '300.0000'),
                (1, '3101', 4, '0', '400.0000'),
                (1, '4201', 5, '0', '275.0000'),
                (1, '4299', 6, '50.0000', '0'),
                (1, '5101', 7, '25.0000', '0'),
                (1, '6101', 8, '120.0000', '0'),
                (1, '6199', 9, '0', '20.0000');

            INSERT INTO journalentry(id, date, concept)
            VALUES (2, '2026-08-10', 'venta futura');
            INSERT INTO journalline(entry_id, account_code, account_id, debit, credit) VALUES
                (2, '1101', 1, '10.0000', '0'),
                (2, '4201', 5, '0', '10.0000');
            """
        )
        conn.commit()
    finally:
        conn.close()

    monkeypatch.setenv("AQORATH_DB", str(db))
    monkeypatch.setattr(accounting_rules, "load_catalog", _catalog)
    return db


def _value_for_label(worksheet, label):
    matches = [
        row[3]
        for row in worksheet.iter_rows(values_only=True)
        if len(row) >= 4 and row[1] == label
    ]
    assert len(matches) == 1, (label, matches)
    return matches[0]


def _database_truth(db):
    conn = sqlite3.connect(str(db))
    try:
        accounts = conn.execute(
            "SELECT id, code, name, nature, origin, parent_id FROM account ORDER BY id"
        ).fetchall()
        entries = conn.execute(
            "SELECT id, date, concept FROM journalentry ORDER BY id"
        ).fetchall()
        lines = conn.execute(
            "SELECT id, entry_id, account_code, account_id, debit, credit FROM journalline ORDER BY id"
        ).fetchall()
        return accounts, entries, lines
    finally:
        conn.close()


def test_formal_bundle_e2e_derives_exact_net_truth_from_canonical_sqlite(canonical_formal_db):
    import aqorath.application as application

    bundle = application.get_financial_statements_bundle(as_of="2026-07-31")

    assert bundle.snapshot.as_of == "2026-07-31"
    assert bundle.income_statement.income.total == Decimal("225.0000")
    assert bundle.income_statement.costs.total == Decimal("25.0000")
    assert bundle.income_statement.expenses.total == Decimal("100.0000")
    assert bundle.income_statement.result == Decimal("100.0000")

    balance = bundle.balance_sheet
    assert balance.assets.total == Decimal("800.0000")
    assert balance.liabilities.total == Decimal("300.0000")
    assert balance.recorded_equity.total == Decimal("400.0000")
    assert balance.current_result == Decimal("100.0000")
    assert balance.total_equity == Decimal("500.0000")
    assert balance.liabilities_and_equity == Decimal("800.0000")
    assert balance.balance_difference == Decimal("0.0000")


def test_formal_xlsx_e2e_preserves_same_sqlite_truth_through_application(canonical_formal_db):
    import aqorath.application as application

    data = application.get_financial_statements_xlsx(as_of="2026-07-31")
    workbook = load_workbook(BytesIO(data), data_only=False)
    try:
        assert workbook.sheetnames == ["Estado de Resultados", "Balance General"]
        income = workbook["Estado de Resultados"]
        balance = workbook["Balance General"]

        assert income["B2"].value == "2026-07-31"
        assert balance["B2"].value == "2026-07-31"
        assert _value_for_label(income, "Total ingresos") == "225.0000"
        assert _value_for_label(income, "Total costos") == "25.0000"
        assert _value_for_label(income, "Total gastos") == "100.0000"
        assert _value_for_label(income, "Resultado del ejercicio") == "100.0000"
        assert _value_for_label(balance, "Total activo") == "800.0000"
        assert _value_for_label(balance, "Total pasivo") == "300.0000"
        assert _value_for_label(balance, "Total patrimonio") == "500.0000"
        assert _value_for_label(balance, "Total pasivo y patrimonio") == "800.0000"
    finally:
        workbook.close()


def test_formal_pdf_e2e_preserves_same_sqlite_truth_through_application(canonical_formal_db):
    import aqorath.application as application

    data = application.get_financial_statements_pdf(as_of="2026-07-31")

    assert data.startswith(b"%PDF-")
    assert data.rstrip().endswith(b"%%EOF")
    assert b"Estado de Resultados" in data
    assert b"Balance General" in data
    assert data.count(b"2026-07-31") >= 2
    for value in (
        b"225.0000",
        b"25.0000",
        b"100.0000",
        b"800.0000",
        b"300.0000",
        b"500.0000",
    ):
        assert value in data


def test_formal_exports_e2e_apply_same_as_of_and_exclude_future_posting(canonical_formal_db):
    import aqorath.application as application

    july = application.get_financial_statements_bundle(as_of="2026-07-31")
    current = application.get_financial_statements_bundle()

    july_balances = {line.account_code: line.ledger_balance for line in july.snapshot.lines}
    current_balances = {line.account_code: line.ledger_balance for line in current.snapshot.lines}

    assert july_balances["1101"] == Decimal("1000.0000")
    assert july_balances["4201"] == Decimal("-275.0000")
    assert current_balances["1101"] == Decimal("1010.0000")
    assert current_balances["4201"] == Decimal("-285.0000")
    assert july.income_statement.result == Decimal("100.0000")
    assert current.income_statement.result == Decimal("110.0000")

    xlsx = application.get_financial_statements_xlsx(as_of="2026-07-31")
    pdf = application.get_financial_statements_pdf(as_of="2026-07-31")
    assert b"285.0000" not in pdf
    workbook = load_workbook(BytesIO(xlsx), data_only=False)
    try:
        values = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
        assert "285.0000" not in values
        assert "1010.0000" not in values
    finally:
        workbook.close()


def test_formal_exports_e2e_are_read_only_against_canonical_sqlite(canonical_formal_db):
    import aqorath.application as application

    before = _database_truth(canonical_formal_db)
    application.get_financial_statements_xlsx(as_of="2026-07-31")
    application.get_financial_statements_pdf(as_of="2026-07-31")
    after = _database_truth(canonical_formal_db)

    assert after == before


def test_formal_exports_e2e_fail_closed_on_unbalanced_sqlite_truth(canonical_formal_db):
    import aqorath.application as application

    conn = sqlite3.connect(str(canonical_formal_db))
    try:
        conn.execute(
            "UPDATE journalline SET debit = '1001.0000' WHERE entry_id = 1 AND account_code = '1101'"
        )
        conn.commit()
    finally:
        conn.close()

    with pytest.raises(ValueError):
        application.get_financial_statements_xlsx(as_of="2026-07-31")
    with pytest.raises(ValueError):
        application.get_financial_statements_pdf(as_of="2026-07-31")
