"""Phase 4F.1 — pure XLSX financial-report renderer contracts."""

from decimal import Decimal
from inspect import signature
from io import BytesIO

import pytest
from openpyxl import load_workbook


def _snapshot(as_of="2026-07-31"):
    from aqorath.reporting import (
        FinancialReportLine,
        FinancialReportSnapshot,
        FinancialReportTotal,
    )

    return FinancialReportSnapshot(
        as_of=as_of,
        lines=(
            FinancialReportLine(
                account_code="1101",
                account_name="Bancos",
                account_type="Activo",
                account_subtype="Circulante",
                nature="DEBIT",
                ledger_balance=Decimal("200.00"),
                normal_balance=Decimal("200.00"),
            ),
            FinancialReportLine(
                account_code="4201",
                account_name="Venta de productos elaborados",
                account_type="Ingreso",
                account_subtype="Actividades con causa",
                nature="CREDIT",
                ledger_balance=Decimal("-200.00"),
                normal_balance=Decimal("200.00"),
            ),
        ),
        totals=(
            FinancialReportTotal("Activo", Decimal("200.00")),
            FinancialReportTotal("Pasivo", Decimal("0")),
            FinancialReportTotal("Patrimonio", Decimal("0")),
            FinancialReportTotal("Ingreso", Decimal("200.00")),
            FinancialReportTotal("Costo", Decimal("0")),
            FinancialReportTotal("Gasto", Decimal("0")),
        ),
        result=Decimal("200.00"),
    )


def _workbook(data):
    return load_workbook(BytesIO(data), data_only=False)


def test_xlsx_renderer_public_contract_and_exact_signature_exist():
    import aqorath.reporting_xlsx as reporting_xlsx

    assert callable(reporting_xlsx.render_financial_report_xlsx)
    params = signature(reporting_xlsx.render_financial_report_xlsx).parameters
    assert list(params) == ["snapshot"]


def test_xlsx_renderer_returns_valid_single_sheet_workbook_with_title_and_as_of():
    import aqorath.reporting_xlsx as reporting_xlsx

    data = reporting_xlsx.render_financial_report_xlsx(_snapshot())
    assert isinstance(data, bytes)
    assert data.startswith(b"PK")

    wb = _workbook(data)
    assert wb.sheetnames == ["Financial Report"]
    ws = wb["Financial Report"]
    assert ws["A1"].value == "Aqorath Financial Report"
    assert ws["A2"].value == "As of"
    assert ws["B2"].value == "2026-07-31"
    assert ws.freeze_panes == "A5"
    wb.close()


def test_xlsx_renderer_preserves_account_rows_in_snapshot_order():
    import aqorath.reporting_xlsx as reporting_xlsx

    wb = _workbook(reporting_xlsx.render_financial_report_xlsx(_snapshot()))
    ws = wb["Financial Report"]

    assert [ws.cell(4, col).value for col in range(1, 8)] == [
        "Account Code",
        "Account Name",
        "Type",
        "Subtype",
        "Nature",
        "Ledger Balance",
        "Normal Balance",
    ]
    assert [ws.cell(5, col).value for col in range(1, 8)] == [
        "1101",
        "Bancos",
        "Activo",
        "Circulante",
        "DEBIT",
        "200.00",
        "200.00",
    ]
    assert [ws.cell(6, col).value for col in range(1, 8)] == [
        "4201",
        "Venta de productos elaborados",
        "Ingreso",
        "Actividades con causa",
        "CREDIT",
        "-200.00",
        "200.00",
    ]
    wb.close()


def test_xlsx_renderer_preserves_snapshot_totals_and_result_without_formulas():
    import aqorath.reporting_xlsx as reporting_xlsx

    wb = _workbook(reporting_xlsx.render_financial_report_xlsx(_snapshot()))
    ws = wb["Financial Report"]

    rows = list(ws.iter_rows(values_only=True))
    total_rows = [row for row in rows if row[0] == "TOTAL"]
    assert [(row[2], row[6]) for row in total_rows] == [
        ("Activo", "200.00"),
        ("Pasivo", "0"),
        ("Patrimonio", "0"),
        ("Ingreso", "200.00"),
        ("Costo", "0"),
        ("Gasto", "0"),
    ]

    result_rows = [row for row in rows if row[0] == "RESULT"]
    assert len(result_rows) == 1
    assert result_rows[0][1] == "Resultado del ejercicio"
    assert result_rows[0][6] == "200.00"

    for row in ws.iter_rows():
        for cell in row:
            assert not (isinstance(cell.value, str) and cell.value.startswith("="))
    wb.close()


def test_xlsx_renderer_preserves_decimal_text_exactly_without_float_conversion():
    from aqorath.reporting import FinancialReportLine, FinancialReportSnapshot, FinancialReportTotal
    import aqorath.reporting_xlsx as reporting_xlsx

    precise = Decimal("12345678901234567890.0100")
    snapshot = FinancialReportSnapshot(
        as_of=None,
        lines=(
            FinancialReportLine(
                "1101",
                "Bancos",
                "Activo",
                "Circulante",
                "DEBIT",
                precise,
                precise,
            ),
        ),
        totals=(FinancialReportTotal("Activo", precise),),
        result=Decimal("0.10"),
    )

    wb = _workbook(reporting_xlsx.render_financial_report_xlsx(snapshot))
    ws = wb["Financial Report"]
    assert ws["F5"].value == "12345678901234567890.0100"
    assert ws["G5"].value == "12345678901234567890.0100"

    rows = list(ws.iter_rows(values_only=True))
    assert [row for row in rows if row[0] == "TOTAL"][0][6] == "12345678901234567890.0100"
    assert [row for row in rows if row[0] == "RESULT"][0][6] == "0.10"
    wb.close()


def test_xlsx_renderer_preserves_none_as_of_and_is_logically_deterministic():
    import aqorath.reporting_xlsx as reporting_xlsx

    snapshot = _snapshot(None)
    first = _workbook(reporting_xlsx.render_financial_report_xlsx(snapshot))
    second = _workbook(reporting_xlsx.render_financial_report_xlsx(snapshot))

    first_ws = first["Financial Report"]
    second_ws = second["Financial Report"]
    assert first_ws["B2"].value is None
    assert tuple(first_ws.values) == tuple(second_ws.values)

    first.close()
    second.close()


def test_xlsx_renderer_does_not_mutate_snapshot():
    import aqorath.reporting_xlsx as reporting_xlsx

    snapshot = _snapshot()
    before = snapshot
    reporting_xlsx.render_financial_report_xlsx(snapshot)
    assert snapshot == before


def test_xlsx_renderer_requires_nominal_financial_report_snapshot():
    import aqorath.reporting_xlsx as reporting_xlsx

    class FakeSnapshot:
        lines = ()
        totals = ()
        result = Decimal("0")
        as_of = None

    with pytest.raises(TypeError):
        reporting_xlsx.render_financial_report_xlsx(FakeSnapshot())


def test_xlsx_renderer_is_pure_and_does_not_query_rebuild_or_write_files(monkeypatch):
    import builtins
    import sqlite3

    import aqorath.accounting_rules as accounting_rules
    import aqorath.reporting as reporting
    import aqorath.reporting_runtime as reporting_runtime
    import aqorath.reporting_source as reporting_source
    import aqorath.reporting_xlsx as reporting_xlsx
    import aqorath.storage as storage

    def forbidden(*args, **kwargs):
        raise AssertionError("XLSX renderer must consume only the supplied snapshot in memory")

    monkeypatch.setattr(builtins, "open", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    monkeypatch.setattr(storage, "get_db_path", forbidden)
    monkeypatch.setattr(accounting_rules, "load_catalog", forbidden)
    monkeypatch.setattr(reporting, "build_financial_report_snapshot", forbidden)
    monkeypatch.setattr(reporting_runtime, "get_financial_report_snapshot", forbidden)
    monkeypatch.setattr(reporting_source, "build_financial_report_snapshot_from_sqlite", forbidden)

    data = reporting_xlsx.render_financial_report_xlsx(_snapshot())
    assert data.startswith(b"PK")
