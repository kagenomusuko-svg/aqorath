"""Phase 4P.1 — pure formal financial-statements XLSX renderer contracts."""

from dataclasses import replace
from decimal import Decimal
from inspect import signature
from io import BytesIO

import pytest
from openpyxl import load_workbook


def _line(code, name, account_type, ledger_balance, *, subtype="", nature="DEBIT", normal_balance=Decimal("9999.9999")):
    from aqorath.reporting import FinancialReportLine

    return FinancialReportLine(
        account_code=code,
        account_name=name,
        account_type=account_type,
        account_subtype=subtype,
        nature=nature,
        ledger_balance=ledger_balance,
        normal_balance=normal_balance,
    )


def _bundle(as_of="2026-07-31"):
    from aqorath.financial_statements import build_financial_statements_bundle
    from aqorath.reporting import FinancialReportSnapshot

    snapshot = FinancialReportSnapshot(
        as_of=as_of,
        lines=(
            _line("1101.001", "BBVA", "Activo", Decimal("1000.0000"), subtype="Circulante"),
            _line("1205", "Depreciación acumulada", "Activo", Decimal("-200.0000"), subtype="No circulante", nature="CREDIT"),
            _line("2101", "Proveedores", "Pasivo", Decimal("-300.0000"), subtype="Corto plazo", nature="CREDIT"),
            _line("3101", "Capital social", "Patrimonio", Decimal("-400.0000"), nature="CREDIT"),
            _line("4201", "Ventas", "Ingreso", Decimal("-275.0000"), subtype="Operación", nature="CREDIT"),
            _line("4299", "Devoluciones sobre ventas", "Ingreso", Decimal("50.0000"), subtype="Operación"),
            _line("5101", "Costo de ventas", "Costo", Decimal("25.0000"), subtype="Operación"),
            _line("6101", "Servicios", "Gasto", Decimal("120.0000"), subtype="Administración"),
            _line("6199", "Recuperación de gasto", "Gasto", Decimal("-20.0000"), subtype="Administración", nature="CREDIT"),
        ),
        totals=(),
        result=Decimal("777777.7777"),
    )
    return build_financial_statements_bundle(snapshot)


def _workbook(bundle):
    from aqorath.financial_statements_xlsx import render_financial_statements_xlsx

    data = render_financial_statements_xlsx(bundle)
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"
    return load_workbook(BytesIO(data), data_only=False)


def _rows_by_code(worksheet):
    result = {}
    for row in worksheet.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0] and row[0][0].isdigit():
            result[row[0]] = row
    return result


def _value_for_label(worksheet, label):
    matches = []
    for row in worksheet.iter_rows(values_only=True):
        if len(row) >= 4 and row[1] == label:
            matches.append(row[3])
    assert len(matches) == 1, (label, matches)
    return matches[0]


def test_formal_xlsx_renderer_public_contract_and_exact_signature_exist():
    import aqorath.financial_statements_xlsx as renderer

    assert callable(renderer.render_financial_statements_xlsx)
    params = signature(renderer.render_financial_statements_xlsx).parameters
    assert list(params) == ["bundle"]


def test_formal_xlsx_renderer_returns_two_named_statement_sheets_with_titles_and_as_of():
    workbook = _workbook(_bundle("2026-07-31"))
    try:
        assert workbook.sheetnames == ["Estado de Resultados", "Balance General"]

        income = workbook["Estado de Resultados"]
        balance = workbook["Balance General"]
        assert income["A1"].value == "Estado de Resultados"
        assert balance["A1"].value == "Balance General"
        assert income["A2"].value == "Al"
        assert balance["A2"].value == "Al"
        assert income["B2"].value == "2026-07-31"
        assert balance["B2"].value == "2026-07-31"
    finally:
        workbook.close()


def test_income_statement_sheet_preserves_line_order_identity_and_presentation_signs():
    workbook = _workbook(_bundle())
    try:
        worksheet = workbook["Estado de Resultados"]
        rows = _rows_by_code(worksheet)
        assert list(rows) == ["4201", "4299", "5101", "6101", "6199"]

        assert rows["4201"][:4] == ("4201", "Ventas", "Operación", "275.0000")
        assert rows["4299"][:4] == ("4299", "Devoluciones sobre ventas", "Operación", "-50.0000")
        assert rows["5101"][:4] == ("5101", "Costo de ventas", "Operación", "25.0000")
        assert rows["6101"][:4] == ("6101", "Servicios", "Administración", "120.0000")
        assert rows["6199"][:4] == ("6199", "Recuperación de gasto", "Administración", "-20.0000")
    finally:
        workbook.close()


def test_income_statement_sheet_uses_view_totals_and_result_verbatim_without_formulas():
    bundle = _bundle()
    income = replace(
        bundle.income_statement,
        income=replace(bundle.income_statement.income, total=Decimal("225.123400")),
        costs=replace(bundle.income_statement.costs, total=Decimal("25.234500")),
        expenses=replace(bundle.income_statement.expenses, total=Decimal("100.345600")),
        result=Decimal("99.543300"),
    )
    sentinel_bundle = replace(bundle, income_statement=income)

    workbook = _workbook(sentinel_bundle)
    try:
        worksheet = workbook["Estado de Resultados"]
        assert _value_for_label(worksheet, "Total ingresos") == "225.123400"
        assert _value_for_label(worksheet, "Total costos") == "25.234500"
        assert _value_for_label(worksheet, "Total gastos") == "100.345600"
        assert _value_for_label(worksheet, "Resultado del ejercicio") == "99.543300"
        assert all(cell.data_type != "f" for row in worksheet.iter_rows() for cell in row)
    finally:
        workbook.close()


def test_balance_sheet_preserves_signed_view_lines_and_uses_view_totals_verbatim():
    bundle = _bundle()
    balance = replace(
        bundle.balance_sheet,
        assets=replace(bundle.balance_sheet.assets, total=Decimal("800.111100")),
        liabilities=replace(bundle.balance_sheet.liabilities, total=Decimal("300.222200")),
        recorded_equity=replace(bundle.balance_sheet.recorded_equity, total=Decimal("400.333300")),
        current_result=Decimal("100.444400"),
        total_equity=Decimal("500.555500"),
        liabilities_and_equity=Decimal("800.777700"),
    )
    sentinel_bundle = replace(bundle, balance_sheet=balance)

    workbook = _workbook(sentinel_bundle)
    try:
        worksheet = workbook["Balance General"]
        rows = _rows_by_code(worksheet)
        assert list(rows) == ["1101.001", "1205", "2101", "3101"]
        assert rows["1101.001"][:4] == ("1101.001", "BBVA", "Circulante", "1000.0000")
        assert rows["1205"][:4] == ("1205", "Depreciación acumulada", "No circulante", "-200.0000")
        assert rows["2101"][:4] == ("2101", "Proveedores", "Corto plazo", "300.0000")
        # XLSX blank cells reload through openpyxl as None, not an empty string.
        assert rows["3101"][:4] == ("3101", "Capital social", None, "400.0000")

        assert _value_for_label(worksheet, "Total activo") == "800.111100"
        assert _value_for_label(worksheet, "Total pasivo") == "300.222200"
        assert _value_for_label(worksheet, "Total patrimonio registrado") == "400.333300"
        assert _value_for_label(worksheet, "Resultado del ejercicio") == "100.444400"
        assert _value_for_label(worksheet, "Total patrimonio") == "500.555500"
        assert _value_for_label(worksheet, "Total pasivo y patrimonio") == "800.777700"
        assert all(cell.data_type != "f" for row in worksheet.iter_rows() for cell in row)
    finally:
        workbook.close()


def test_formal_xlsx_renderer_preserves_decimal_text_exactly_without_float_cells():
    workbook = _workbook(_bundle())
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            monetary_values = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.column == 4 and cell.value is not None:
                        monetary_values.append(cell)
            assert monetary_values
            for cell in monetary_values:
                if cell.value == "Importe":
                    continue
                assert isinstance(cell.value, str)
                assert cell.data_type != "n"
                assert cell.data_type != "f"
    finally:
        workbook.close()


def test_formal_xlsx_renderer_preserves_none_as_of_as_blank_on_both_sheets():
    workbook = _workbook(_bundle(as_of=None))
    try:
        assert workbook["Estado de Resultados"]["B2"].value is None
        assert workbook["Balance General"]["B2"].value is None
    finally:
        workbook.close()


def test_formal_xlsx_renderer_requires_nominal_financial_statements_bundle():
    from aqorath.financial_statements_xlsx import render_financial_statements_xlsx

    class FakeBundle:
        pass

    with pytest.raises(TypeError):
        render_financial_statements_xlsx(FakeBundle())


def test_formal_xlsx_renderer_is_logically_deterministic_and_does_not_mutate_bundle():
    bundle = _bundle()
    before = bundle

    first = _workbook(bundle)
    second = _workbook(bundle)
    try:
        assert bundle == before
        for sheet_name in first.sheetnames:
            first_rows = list(first[sheet_name].iter_rows(values_only=True))
            second_rows = list(second[sheet_name].iter_rows(values_only=True))
            assert first_rows == second_rows
    finally:
        first.close()
        second.close()


def test_formal_xlsx_renderer_is_pure_uses_bundle_only_and_never_writes_files(monkeypatch):
    import builtins
    import sqlite3
    import openpyxl.worksheet._writer as worksheet_writer

    import aqorath.accounting_rules as accounting_rules
    import aqorath.balance_sheet as balance_sheet
    import aqorath.financial_statement_semantics as semantics
    import aqorath.financial_statements as financial_statements
    import aqorath.financial_statements_xlsx as renderer
    import aqorath.income_statement as income_statement
    import aqorath.reporting_runtime as reporting_runtime
    import aqorath.reporting_xlsx as reporting_xlsx
    import aqorath.storage as storage

    bundle = _bundle()

    def forbidden(*args, **kwargs):
        raise AssertionError("formal XLSX renderer must consume only its supplied bundle")

    monkeypatch.setattr(builtins, "open", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    monkeypatch.setattr(worksheet_writer, "create_temporary_file", forbidden)
    monkeypatch.setattr(storage, "get_db_path", forbidden)
    monkeypatch.setattr(accounting_rules, "load_catalog", forbidden)
    monkeypatch.setattr(reporting_runtime, "get_financial_report_snapshot", forbidden)
    monkeypatch.setattr(reporting_runtime, "get_financial_statements_bundle", forbidden)
    monkeypatch.setattr(financial_statements, "build_financial_statements_bundle", forbidden)
    monkeypatch.setattr(income_statement, "build_income_statement_view", forbidden)
    monkeypatch.setattr(balance_sheet, "build_balance_sheet_view", forbidden)
    monkeypatch.setattr(semantics, "compute_financial_statement_totals", forbidden)
    monkeypatch.setattr(reporting_xlsx, "render_financial_report_xlsx", forbidden)

    data = renderer.render_financial_statements_xlsx(bundle)
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"
